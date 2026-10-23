"""ARCHIVED: legacy checkpoint report generator, not the full30 v2 report.

Build five Chinese slides from measured inference/ablation summaries only.

The batch summary's sibling <video_id>/diagnostics.json files are read only for
frame indices/count, strict loading, temporal mode and seed metadata. Neither
source annotations nor ground-truth target values are opened by this script.
"""

import argparse
import json
import math
import os
from pathlib import Path
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OFFICIAL_IDS = ("000237", "001483", "003235", "006543", "024386",
                "033447", "035729", "036120", "047771", "051657")
NAVY = "172C46"
TEAL = "087D83"
GOLD = "A86D12"
INK = "26374B"
MUTED = "5A6D80"
PALE = "F3F6FA"
LINE = "DCE4EC"
WHITE = "FFFFFF"
AMBER_PALE = "FFF3DC"
FONT = "Microsoft YaHei"
REPORT_TITLE = "VLSOT基线复现与提交材料（待质量核验）"
LOCALIZATION_WARNING = "定位误差异常，当前权重/训练编码仍待核对；不能据此宣称效果可靠。"


def _read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _number(value, name, nonnegative=True):
    if isinstance(value, bool):
        raise ValueError("{} must be numeric".format(name))
    value = float(value)
    if not math.isfinite(value) or (nonnegative and value < 0):
        raise ValueError("{} must be finite and nonnegative".format(name))
    return value


def collect_data(batch_summary, ablation, onnx_manifest=None):
    batch_path = Path(batch_summary).resolve()
    batch = _read(batch_path)
    clips = batch.get("clips")
    if not isinstance(clips, list) or len(clips) != 10:
        raise ValueError("Report requires ten completed clip summaries")
    identifiers = [str(clip.get("video_id", "")) for clip in clips]
    if len(set(identifiers)) != 10 or set(identifiers) != set(OFFICIAL_IDS):
        raise ValueError("Batch IDs must be exactly the fixed ten official IDs")
    for clip in clips:
        _number(clip["elapsed_seconds"], "clip elapsed_seconds")
        _number(clip["peak_cuda_memory_mib"], "peak_cuda_memory_mib")
    seeds = set()
    modes = set()
    total_frames = 0
    for video_id in OFFICIAL_IDS:
        diagnostic = _read(batch_path.parent / video_id / "diagnostics.json")
        if str(diagnostic.get("video_id")) != video_id:
            raise ValueError("Sibling diagnostics video ID mismatch")
        frames = diagnostic.get("frames")
        if not isinstance(frames, list) or len(frames) != 30:
            raise ValueError("{} has not completed 30 frames".format(video_id))
        indexes = [frame.get("frame_index") for frame in frames]
        if (any(type(index) is not int for index in indexes) or
                sorted(indexes) != list(range(1, 31))):
            raise ValueError("{} requires unique frame indices 1..30".format(video_id))
        if diagnostic.get("state_dict_load") != "strict_success":
            raise ValueError("{} lacks strict task-weight loading evidence".format(video_id))
        seeds.add(diagnostic.get("seed_per_video"))
        modes.add(diagnostic.get("temporal_mode"))
        total_frames += len(indexes)
    if total_frames != 300:
        raise ValueError("Expected exactly 300 completed frame records")
    if len(modes) != 1 or None in modes or next(iter(modes)) != batch.get("temporal_mode"):
        raise ValueError("Clip temporal modes disagree with the batch summary")
    if len(seeds) != 1 or None in seeds:
        raise ValueError("Clip seed metadata is missing or inconsistent")

    comparison = _read(ablation)
    if str(comparison.get("video_id")) != "039373":
        raise ValueError("This report's diagnostic is restricted to training clip 039373")
    if comparison.get("checkpoint_identical") is not True:
        raise ValueError("Ablation must record use of identical task weights")
    if comparison.get("seed_per_video") is None:
        raise ValueError("Ablation must record its seed")
    records = comparison.get("results")
    if not isinstance(records, list) or len(records) != 2:
        raise ValueError("Ablation requires exactly carry and reset_each_frame records")
    by_mode = {record.get("mode"): record for record in records}
    if set(by_mode) != {"carry", "reset_each_frame"}:
        raise ValueError("Ablation modes must be carry and reset_each_frame")
    for record in records:
        if record.get("frames") != 30:
            raise ValueError("Training diagnostic must contain 30 frames per mode")
        for key in ("mean_center_distance_m", "median_center_distance_m", "forward_mean_seconds"):
            _number(record[key], key)
        if len(record.get("dimension_mae_lwh_m", [])) != 3:
            raise ValueError("Dimension MAE must be in l,w,h order")
        for value in record["dimension_mae_lwh_m"]:
            _number(value, "dimension MAE")

    return {
        "batch": batch,
        "batch_source": str(batch_path),
        "ablation": comparison,
        "ablation_by_mode": by_mode,
        "onnx": _read(onnx_manifest) if onnx_manifest else None,
        "total_frames": total_frames,
        "total_seconds": _number(batch["total_seconds"], "total_seconds"),
        "peak_memory_gib": max(float(clip["peak_cuda_memory_mib"]) for clip in clips) / 1024.,
        "seed": next(iter(seeds)),
        "temporal_mode": next(iter(modes)),
        "all_dimensions_positive": all(clip.get("raw_dimensions_all_positive") is True for clip in clips),
        "localization_anomaly": any(float(record["mean_center_distance_m"]) > 100
                                    for record in records),
    }


def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


def _font(run, size, color=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    properties = run._r.get_or_add_rPr()
    properties.set("lang", "zh-CN")
    east_asian = properties.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.append(east_asian)
    east_asian.set("typeface", FONT)


def _shape(slide, x, y, w, h, fill, border=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(border or fill)
    if rounded:
        shape.adjustments[0] = .10
    return shape


def _text(slide, x, y, w, h, text, size=20, color=INK, bold=False,
          align=PP_ALIGN.LEFT, line_space=1.13):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.015)
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(8)
        paragraph.line_spacing = line_space
        run = paragraph.add_run()
        run.text = line
        _font(run, size, color, bold)
    return shape


def _slide(presentation, title, subtitle, page):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(PALE)
    _shape(slide, 0, 0, 13.333333, .12, TEAL)
    _text(slide, .58, .43, 12.15, .62, title, 30, NAVY, True)
    _text(slide, .60, 1.10, 12.0, .48, subtitle, 17.5, MUTED)
    _shape(slide, .60, 7.08, 12.12, .012, LINE)
    _text(slide, .60, 7.15, 11.2, .20,
          "长安大学第二届人工智能创新大赛 · 专项命题挑战赛道 / VLSOT", 10.5, MUTED)
    _text(slide, 11.9, 7.13, .8, .23, "{} / 5".format(page), 11, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def _card(slide, x, y, w, h, title, body, body_size=19):
    _shape(slide, x, y, w, h, WHITE, LINE, True)
    _shape(slide, x + .20, y + .22, .06, .34, TEAL)
    _text(slide, x + .38, y + .19, w - .60, .50, title, 21, NAVY, True)
    _text(slide, x + .25, y + .86, w - .50, h - 1.03, body, body_size)


def _stat(slide, x, label, value, note):
    _shape(slide, x, 1.84, 3.89, 1.85, WHITE, LINE, True)
    _text(slide, x + .21, 2.00, 3.47, .35, label, 18, MUTED)
    _text(slide, x + .21, 2.40, 3.47, .58, value, 32, TEAL, True)
    _text(slide, x + .21, 3.13, 3.47, .30, note, 17, MUTED)


def _time(seconds):
    return "{} 分 {:.1f} 秒".format(int(seconds // 60), seconds % 60)


def _onnx_text(manifest):
    if manifest is None:
        return "未提供 ONNX 验证记录", "运行验证状态待核验。\n报告未据此认定 ONNX Runtime 已成功运行。"
    status = str(manifest.get("status", "未记录"))
    checks = manifest.get("two_step_parity")
    verified = (manifest.get("onnx_runtime_verified") is True and status == "validated"
                and isinstance(checks, list) and len(checks) == 2
                and manifest.get("graph", {}).get("checker") == "passed")
    if verified:
        provider = manifest.get("ort_provider", "未记录")
        if isinstance(provider, list):
            provider = " / ".join(provider)
        lines = ["Checker 与 ORT 两步状态回馈对比通过。",
                 "执行后端：{}".format(provider),
                 "容差：rtol={}，atol={}".format(manifest.get("parity_rtol"), manifest.get("parity_atol"))]
        reference = manifest.get("ttm_original_reference_checks", [])
        if len(reference) == 3 and all(check.get("passed") is True for check in reference):
            lines.append("TTM 原实现与状态适配器：3 个相位核对通过。")
        return "ONNX 实测验证通过", "\n".join(lines)
    error = " ".join(str(manifest.get("error", "尚无完整成功验证证据")).split())
    if len(error) > 150:
        error = error[:147] + "…"
    return "ONNX 当前状态：{}".format(status), "{}\n该记录未证明 ORT 完整推理成功。".format(error)


def build_presentation(data):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = REPORT_TITLE
    presentation.core_properties.subject = "既有模型复现、推理适配与真实运行诊断"
    presentation.core_properties.author = ""
    batch = data["batch"]
    localization_anomaly = any(
        float(record["mean_center_distance_m"]) > 100
        for record in data["ablation_by_mode"].values())

    slide = _slide(presentation, REPORT_TITLE,
                   "任务与基线：以视频帧和目标文字描述为输入，输出每帧三维尺寸与位置", 1)
    _card(slide, .60, 1.86, 7.48, 3.04, "继承的技术基础",
          "• Mono3DVG 视觉语言框架\n• Swin-Large 图像特征与 RoBERTa 文本编码\n• 既有项目中的 TTM 时序记忆模块", 21)
    _card(slide, 8.30, 1.86, 4.42, 3.04, "权重与贡献边界",
          "继承师兄提供的已有任务权重。\n本次工作聚焦流程复现、推理适配与提交材料。\n网络结构及既有权重不列为原创。", 19)
    _text(slide, .68, 5.18, 12.0, .43, "队名：________________    队长：________________    学院：________________", 20, NAVY)
    _text(slide, .68, 5.92, 12.0, .70,
          "基线来源：项目源码中的 Mono3DVG / Swin / RoBERTa / TTM 实现\n本报告所有数值来自本次保存的运行汇总；官方成绩以 VLP 平台评测为准。", 18, MUTED)

    slide = _slide(presentation, "方法与本次工程贡献",
                   "沿用既有模型，补齐跨视频状态隔离、原生解码、坐标转换和提交校验", 2)
    _card(slide, .60, 1.80, 5.96, 2.24, "01  时序状态隔离",
          "每段视频开始清空 TTM 记忆与计数器。\n片内按运行模式保留状态，避免跨视频串扰。", 19)
    _card(slide, 6.78, 1.80, 5.94, 2.24, "02  按原标签编码解码",
          "归一化 cx/cy → 输入像素 → 逆 affine。\n第二轴加回预测 h/2；尺寸 h/w/l 转 l/w/h。", 19)
    _card(slide, .60, 4.28, 5.96, 2.24, "03  来源核查与坐标转换",
          "原生监督来源核查为路端虚拟 LiDAR。\n逐帧使用官方外参 R·p+t 转世界坐标，并保留原生备份。", 19)
    _card(slide, 6.78, 4.28, 5.94, 2.24, "04  提交结构与数值校验",
          "固定 10 个 ID；每个 30 行、每行 7 列。\n校验有限值、正尺寸与 result/ 层级；正式推理不读测试目标数值。", 19)
    coordinate_source = _text(slide, .71, 6.67, 11.96, .30,
                              "坐标来源：AIR-THU / DAIR-V2X-Seq 官方数据说明", 17, MUTED)
    coordinate_source.text_frame.paragraphs[0].runs[0].hyperlink.address = (
        "https://github.com/AIR-THU/DAIR-V2X-Seq/blob/main/dataset/v2x-seq-spd/README.md")

    slide = _slide(presentation, "正式推理：10 个片段，300 帧完成",
                   "已逐一检查固定序列 ID、每段帧号 1–30 与 strict 权重加载记录", 3)
    _stat(slide, .60, "完整性核验", "10 段 / 300 帧", "每段 30 个唯一帧索引")
    _stat(slide, 4.72, "实测总用时", _time(data["total_seconds"]), "含模型加载、预处理与 IO")
    _stat(slide, 8.83, "实测峰值显存", "{:.2f} GiB".format(data["peak_memory_gib"]), "取各片段峰值中的最大值")
    _text(slide, .69, 4.02, 12.0, .60,
          "设备：{}    |    PyTorch：{}".format(batch.get("gpu", "未记录"), batch.get("torch", "未记录")), 19, NAVY)
    _text(slide, .69, 4.67, 12.0, .74,
          "序列：" + " · ".join(OFFICIAL_IDS[:5]) + "\n        " + " · ".join(OFFICIAL_IDS[5:]), 18, MUTED)
    dimension_note = "汇总记录：全部原生尺寸为正。" if data["all_dimensions_positive"] else "汇总记录中存在尺寸检查异常或缺失，需复核。"
    _text(slide, .69, 5.54, 12.0, .42, dimension_note, 18, MUTED)
    _shape(slide, .68, 6.10, 11.98, .62, AMBER_PALE, rounded=True)
    _text(slide, .88, 6.22, 11.5, .36, "官方评测分数：待 VLP 平台返回；本地完整性通过不等于模型精度已确认。", 19, GOLD, True)

    slide = _slide(presentation, "TTM 运行消融：保留记忆与逐帧重置",
                   "仅使用真实训练序列 039373；读取已计算的消融汇总，不读取原始标签", 4)
    _shape(slide, .61, 1.75, 12.10, .72, AMBER_PALE, rounded=True)
    _text(slide, .82, 1.92, 11.65, .40,
          "单个训练片段 · 30 个相关帧 · 运行诊断；不是验证集成绩，不能推断泛化。", 20, GOLD, True)
    table = slide.shapes.add_table(3, 5, Inches(.65), Inches(2.79), Inches(12.02), Inches(1.96)).table
    widths = (2.15, 2.20, 2.08, 3.25, 2.34)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    headers = ("时序模式", "中心距离均值 / m", "中位数 / m", "尺寸 MAE l/w/h / m", "单帧前向 / s")
    rows = [headers]
    for mode, label in (("carry", "carry\n跨帧保留"), ("reset_each_frame", "reset\n逐帧重置")):
        record = data["ablation_by_mode"][mode]
        rows.append((label, "{:.4f}".format(record["mean_center_distance_m"]),
                     "{:.4f}".format(record["median_center_distance_m"]),
                     " / ".join("{:.4f}".format(value) for value in record["dimension_mae_lwh_m"]),
                     "{:.4f}".format(record["forward_mean_seconds"])))
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.07)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(NAVY if row_index == 0 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    _font(run, 17 if row_index == 0 else 18,
                          WHITE if row_index == 0 else INK, row_index == 0)
    comparison = data["ablation"]
    pairing = "相同权重" if comparison.get("checkpoint_identical") is True else "权重一致性未核验"
    _text(slide, .71, 5.05, 11.98, .52,
          "实验条件：{}；每段固定随机种子 {}；仅改变记忆保留策略。".format(pairing, comparison.get("seed_per_video", "未记录")), 19, NAVY)
    if localization_anomaly:
        _shape(slide, .65, 5.70, 12.02, 1.06, AMBER_PALE, rounded=True)
        _text(slide, .84, 5.87, 11.64, .76,
              LOCALIZATION_WARNING + "\n单训练片段的数值差异只作运行诊断，不代表模型提升或平台成绩。", 18, GOLD)
    else:
        _text(slide, .71, 5.78, 11.98, .87,
              "指标为预测与训练源标签 XYZ 的欧氏距离，以及 l/w/h 的平均绝对误差。\n这些数值只描述该训练片段的运行现象，不充当本次平台分数。", 18.5, MUTED)

    slide = _slide(presentation, "复现、ONNX 验证状态与局限",
                   "分别说明已完成的 PyTorch 流程、ONNX 实测证据与尚未解决的评估边界", 5)
    epoch = batch.get("checkpoint_epoch", "未记录")
    mode = "片内保留记忆" if data["temporal_mode"] == "carry" else "逐帧重置记忆"
    _card(slide, .60, 1.80, 5.96, 3.59, "PyTorch 复现路径",
          ("• 继承任务 checkpoint，epoch={}\n• 每段随机种子 {}；{}\n"
           "• batch_infer.py → export_vlsot.py\n• 正式输入仅图像、描述和标定").format(
               epoch, data["seed"], mode), 18.5)
    onnx_title, onnx_body = _onnx_text(data["onnx"])
    _card(slide, 6.78, 1.80, 5.94, 3.59, onnx_title, onnx_body, 18)
    _shape(slide, .65, 5.64, 12.02, 1.06,
           AMBER_PALE if localization_anomaly else WHITE,
           None if localization_anomaly else LINE, True)
    limitations = (
        LOCALIZATION_WARNING + "\n另：权重训练来源与数据独立性尚未完整复核；平台分数及坐标实际解释待确认。"
        if localization_anomaly else
        "局限：单训练片段消融不代表泛化；继承权重的完整训练过程与数据独立性尚未全部复核。\n世界坐标按平台字面规范输出，平台实际解释及最终精度仍待官方评测确认。")
    _text(slide, .85, 5.82, 11.62, .72, limitations, 17.5,
          GOLD if localization_anomaly else MUTED)

    if len(presentation.slides) != 5:
        raise AssertionError("The report must contain exactly five slides")
    return presentation


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch-summary", required=True)
    parser.add_argument("--ablation", required=True)
    parser.add_argument("--onnx-manifest")
    parser.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    target = Path(args.output).resolve()
    if target.suffix.lower() != ".pptx":
        raise ValueError("Output must end in .pptx")
    if target.exists():
        raise FileExistsError("Refusing to overwrite {}".format(target))
    data = collect_data(args.batch_summary, args.ablation, args.onnx_manifest)
    presentation = build_presentation(data)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, stage_name = tempfile.mkstemp(prefix=".vlsot_report_", suffix=".pptx", dir=str(target.parent))
    os.close(handle)
    stage = Path(stage_name)
    try:
        presentation.save(str(stage))
        reloaded = Presentation(str(stage))
        if len(reloaded.slides) != 5:
            raise AssertionError("Saved report does not contain exactly five slides")
        os.link(str(stage), str(target))  # Atomic publication without overwrites.
    finally:
        stage.unlink()
    print(json.dumps({"output": str(target), "slides": 5,
                      "verified_clips": 10, "verified_frames": data["total_frames"],
                      "onnx_manifest_supplied": data["onnx"] is not None,
                      "ground_truth_files_read": False}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
